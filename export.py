"""Выгрузка в Excel: каталог, сводка по категориям и изменения с прошлого запуска."""

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CHANGE_FILLS = {
    "New": PatternFill("solid", fgColor="E2EFDA"),
    "Removed": PatternFill("solid", fgColor="FCE4D6"),
    "Price down": PatternFill("solid", fgColor="DDEBF7"),
    "Price up": PatternFill("solid", fgColor="FFF2CC"),
    "Back in stock": PatternFill("solid", fgColor="E2EFDA"),
    "Out of stock": PatternFill("solid", fgColor="FCE4D6"),
}
PRICE_FORMAT = "£#,##0.00"


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _finish(ws, price_cols: list[str]):
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in price_cols:
        for cell in ws[col][1:]:
            cell.number_format = PRICE_FORMAT


def _catalog(wb, products):
    ws = _sheet(wb, "Catalog", ["Title", "Category", "Price", "Rating", "In stock", "Link"],
                [55, 20, 11, 10, 10, 14])
    for p in sorted(products, key=lambda p: (p["category"], p["title"])):
        ws.append([p["title"], p["category"], p["price"], "★" * p["rating"],
                   "Yes" if p["in_stock"] else "No", "Open"])
        link = ws.cell(row=ws.max_row, column=6)
        link.hyperlink, link.style = p["url"], "Hyperlink"
    _finish(ws, ["C"])
    # Цветовая шкала цен: дешёвые — зелёные, дорогие — красные
    ws.conditional_formatting.add(
        f"C2:C{ws.max_row}",
        ColorScaleRule(start_type="min", start_color="63BE7B",
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color="F8696B"),
    )


def _summary(wb, products):
    ws = _sheet(wb, "Summary", ["Category", "Products", "Avg price", "Min price", "Max price", "Avg rating"],
                [22, 11, 12, 12, 12, 12])
    groups = defaultdict(list)
    for p in products:
        groups[p["category"]].append(p)
    for category, items in sorted(groups.items(), key=lambda g: -len(g[1])):
        prices = [p["price"] for p in items]
        ws.append([category, len(items), round(sum(prices) / len(prices), 2), min(prices), max(prices),
                   round(sum(p["rating"] for p in items) / len(items), 1)])
    _finish(ws, ["C", "D", "E"])


def _changes(wb, changes):
    ws = _sheet(wb, "Changes", ["Change", "Title", "Category", "Old price", "New price", "Diff", "Link"],
                [15, 55, 20, 11, 11, 10, 14])
    for c in changes:
        ws.append([c["type"], c["title"], c["category"], c.get("old_price"), c.get("new_price"),
                   c.get("diff"), "Open"])
        row = ws.max_row
        ws.cell(row=row, column=1).fill = CHANGE_FILLS.get(c["type"], PatternFill())
        link = ws.cell(row=row, column=7)
        link.hyperlink, link.style = c["url"], "Hyperlink"
    if not changes:
        ws.append(["No changes since the previous run"])
    _finish(ws, ["D", "E", "F"])


def save_report(products: list[dict], changes: list[dict], path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _changes(wb, changes)
    _catalog(wb, products)
    _summary(wb, products)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
